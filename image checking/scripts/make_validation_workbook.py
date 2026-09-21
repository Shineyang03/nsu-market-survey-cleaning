"""make_validation_workbook.py -- the workbook a person reads to establish ground truth.

WHAT IT BUILDS. One .xlsx with a photograph embedded per row and empty columns for the
reader's answer. That is all. It carries no typed weight, no published value, no model
reading, no item name, no stratum.

WHY NOTHING IS PRE-FILLED. `docs/adjudication_playbook.md` says to pre-fill every
verdict, because confirm/override is cheap and decide-from-scratch is not. That rule is
right when a reviewer is ADJUDICATING a judgement call. It is wrong here. The human is
not adjudicating anything -- they are the measuring instrument, the only non-circular
ground truth this project has for what a scale display said. Show them "the model read
0.095" and they will confirm it, inheriting the exact error the exercise exists to
detect. The playbook's efficiency cost is real and is accepted.

For the same reason the workbook does not say which images the models disagreed on. A
reader who knows a row is contentious reads it differently.

WHY PYTHON. .xlsx writing is one of the three things the project's CLAUDE.md reserves
for Python. The draw itself is Stata (`07_human_validation_draw.do`); this script
chooses nothing.

THE IMAGES ARE EMBEDDED, not linked. A workbook of hyperlinks into a Box path breaks the
moment it is emailed, moved, or opened on a machine that syncs to a different letter --
and a validation set that cannot be opened is a validation set that does not get done.

USAGE
    python make_validation_workbook.py --ids ../outputs/tables/human_validation_ids.csv \
        --out ../outputs/qc/human_validation_v1.xlsx
"""

from __future__ import annotations

import argparse
import os
import sys

import pandas as pd
from PIL import Image, ImageOps
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
IMGCHK = os.path.dirname(HERE)
BRIDGE = os.path.join(IMGCHK, "outputs", "bridge", "photo_id_bridge.csv")

# Columns the reader fills. Everything else is locked context.
ANSWER_COLS = [
    ("display_text", 18,
     "EXACTLY as shown, e.g. 0.800 -- keep the decimal point and all zeros"),
    ("legible", 14, "clear / probable / ambiguous / not_visible"),
    ("scale_present", 13, "yes / no"),
    ("notes", 40, "anything that qualifies your answer"),
]


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--ids", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--px", type=int, default=820,
                    help="embedded image width; big enough to read a display")
    args = ap.parse_args(argv)

    ids = pd.read_csv(args.ids)
    bridge = pd.read_csv(BRIDGE)[["id", "filename", "photo_path"]]
    df = ids[["id"]].merge(bridge, on="id", how="left")
    if df.photo_path.isna().any():
        print("some ids have no photograph in the bridge", file=sys.stderr)
        return 2

    # SHUFFLED, seeded. The draw is sorted by id, and id is assigned on a content key
    # starting with province -- so an unshuffled workbook walks the reader through one
    # market at a time, priming them.
    df = df.sample(frac=1.0, random_state=20260920).reset_index(drop=True)
    df["seq"] = range(1, len(df) + 1)

    # THE JOIN KEY IS THE IMAGE CODE, NOT THE ROW NUMBER. `seq` is a position: it
    # changes if the set is re-drawn, re-shuffled, re-sent, or if someone sorts the
    # sheet, and a returned workbook whose positions have moved cannot be attached to
    # anything. The image code is the crosswalk's own key, unique per photograph, and it
    # travels with the row wherever it goes.
    #
    # It is also safe to show. The code is a capture timestamp -- it says nothing about
    # the weight, the item, or what anyone typed. `id` would be weaker on that count:
    # ids are assigned in order of province, municipality and item, so the number itself
    # leaks a little about what the photograph shows.
    df["image_code"] = df.filename.str.replace(".jpg", "", regex=False)

    tmpdir = os.path.join(IMGCHK, "outputs", "qc", "_wb_tmp")
    os.makedirs(tmpdir, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "read the display"

    hdr = ["row", "image_code", "photograph"] + [c for c, _, _ in ANSWER_COLS]
    ws.append(hdr)
    for i, h in enumerate(hdr, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="333333")
        c.alignment = Alignment(vertical="center", horizontal="center")
    ws.freeze_panes = "A2"

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = int(args.px / 7.2)
    for i, (name, width, _) in enumerate(ANSWER_COLS, start=4):
        ws.column_dimensions[get_column_letter(i)].width = width

    # A closed vocabulary on `legible` and `scale_present`: free text there would have to
    # be normalised afterwards, and normalising a ground-truth column by hand is how a
    # ground truth stops being one.
    dv_leg = DataValidation(type="list",
                            formula1='"clear,probable,ambiguous,not_visible"',
                            allow_blank=True, showDropDown=False)
    dv_yn = DataValidation(type="list", formula1='"yes,no"',
                           allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv_leg)
    ws.add_data_validation(dv_yn)

    for _, r in df.iterrows():
        row = int(r["seq"]) + 1
        im = ImageOps.exif_transpose(Image.open(r["photo_path"])).convert("RGB")
        im.thumbnail((args.px, args.px), Image.LANCZOS)
        tmp = os.path.join(tmpdir, f"v_{int(r['seq']):03d}.jpg")
        im.save(tmp, quality=90)

        ws.cell(row=row, column=1, value=int(r["seq"])).alignment = Alignment(
            vertical="top", horizontal="center")
        # Text, not a number. These codes are 13 digits; Excel renders a bare integer
        # that long in scientific notation and then round-trips it back as 1.7743E+12,
        # silently destroying the join key.
        kc = ws.cell(row=row, column=2, value=str(r["image_code"]))
        kc.number_format = "@"
        kc.alignment = Alignment(vertical="top", horizontal="left")
        ws.row_dimensions[row].height = im.height * 0.76  # px -> points
        ws.add_image(XLImage(tmp), f"C{row}")

        # display_text IS TEXT, and this was missed in v1 at real cost. Left as a
        # General cell, Excel coerces the answer to a number the moment it is typed:
        # a display read as `0.300` comes back as `0.3`, and the trailing zeros the
        # whole instrument is built to preserve are gone before the file is saved.
        # 37 of 40 v1 answers returned as floats. The numeric comparison survived, so
        # nothing was lost that mattered, but the exact-string check the instrument
        # promises was simply unavailable.
        dc = ws.cell(row=row, column=4)
        dc.number_format = "@"

        dv_leg.add(ws.cell(row=row, column=5))
        dv_yn.add(ws.cell(row=row, column=6))
        for col in (4, 5, 6, 7):
            ws.cell(row=row, column=col).alignment = Alignment(
                vertical="top", horizontal="left", wrap_text=True)

    # ---- the instructions, as their own tab ------------------------------------
    gs = wb.create_sheet("how to read these")
    lines = [
        ("What this is", True),
        ("40 photographs from the NSU market survey. For each one, write what the "
         "scale display shows.", False),
        ("Your readings are the ground truth this project checks its data against. "
         "Nothing you need is hidden from you, and nothing you should not see is "
         "shown: no typed value, no computed value, no model's guess.", False),
        ("", False),
        ("How to fill it in", True),
        ("display_text -- the characters on the display, EXACTLY as shown.", False),
        ("   Keep the decimal point and every zero. 0.800 is not 0.8 and not 800.", False),
        ("   Telling those apart is the whole point of this exercise.", False),
        ("   Leave BLANK if there is no scale, or the display is off or unreadable.", False),
        ("", False),
        ("legible -- how sure you are. 'ambiguous' is a correct and useful answer.", False),
        ("   A wrong number is far worse than saying you could not read it.", False),
        ("", False),
        ("scale_present -- is there an electronic scale in the photograph at all?", False),
        ("   Many of these are photographs of a package label with no scale.", False),
        ("", False),
        ("notes -- anything that qualifies your answer: glare, an angle, a partial", False),
        ("   reading, a second display in shot.", False),
        ("", False),
        ("Two things that matter", True),
        ("These scales usually show THREE stacked red rows: WEIGHT on top, then unit", False),
        ("   price, then total price. Read the WEIGHT row only.", False),
        ("Unlit segments still glow faintly. A row showing a dim 8.8.8.8.8 is BLANK,", False),
        ("   not a row of eights.", False),
        ("", False),
        ("Please do not look up what these items weigh, and do not reason from the", False),
        ("   object to a plausible weight. Read only what the display shows.", False),
    ]
    for i, (txt, bold) in enumerate(lines, start=1):
        c = gs.cell(row=i, column=1, value=txt)
        if bold:
            c.font = Font(bold=True, size=12)
    gs.column_dimensions["A"].width = 100

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    wb.save(args.out)

    # An audit record, NOT a key the returned workbook depends on. The workbook carries
    # its own join key in `image_code`, so answers attach even if this file is lost, the
    # rows are sorted, or the set is re-drawn. This maps the code to the weighing id, so
    # a reader of the outputs can see which weighing each answer belongs to without
    # re-deriving it through the bridge.
    keypath = args.out.replace(".xlsx", "_key.csv")
    df[["image_code", "id", "filename", "seq"]].to_csv(keypath, index=False)

    print(f"images          {len(df)}")
    print(f"embedded at     {args.px}px")
    print(f"workbook        {args.out}")
    print(f"code -> id map  {keypath}")
    print("\nThe workbook carries a photograph and its image code. No typed weight, no")
    print("published value, no model reading, no item name, no stratum.")
    print("The image code is the join key: it is unique per photograph and survives")
    print("sorting, re-drawing and re-sending, which a row number does not.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
