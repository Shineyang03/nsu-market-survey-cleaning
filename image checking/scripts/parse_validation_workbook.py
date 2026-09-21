"""parse_validation_workbook.py -- read the returned human validation workbook.

WHAT IT DOES. Extracts the human's answers, normalises them, and writes a CSV keyed on
`image_code` -- the workbook's own join key, which survives sorting and re-sending.

THE HUMAN'S ANSWERS ARE NEVER REWRITTEN. `display_raw` holds exactly what the cell
contained; every derived column sits beside it. A ground truth that has been tidied in
place is no longer a ground truth, and the project rule is the same one that applies to
the deliverables: never overwrite a raw value with a cleaned one.

THREE THINGS THE RETURNED WORKBOOK REVEALED, all handled here rather than silently:

1. EXCEL COERCED THE READINGS TO NUMBERS. 37 of 40 answers came back as floats, so a
   display read as `0.300` is stored as `0.3`. That is the trailing-zero destruction the
   instrument warns about, and it is a defect in the WORKBOOK GENERATOR -- `image_code`
   was given a text format and `display_text` was not. It is fixed for v2.

   It is survivable here because the distinction that matters -- 0.3 against 3 against
   300 -- is preserved by the number. What is lost is literal character fidelity, so
   `display_exact_ok` records that an exact-string comparison is NOT available for these
   rows and no such comparison is attempted.

2. THE `legible` DROPDOWN DID NOT BIND. Answers came back as free text -- "Y",
   "Y (bad glare)", "Y (just about)". The closed vocabulary was not enforced. The free
   text is in fact MORE informative than the vocabulary would have been, so it is
   parsed rather than rejected: `legible_yes` and `glare_flag`.

3. THREE ROWS RECORD A PACKAGE VOLUME, NOT THE SCALE DISPLAY -- "355 mL", "800 mL". The
   reviewer stated the rule: where both a package volume and a scale reading are
   visible, prefer the volume. Those rows are flagged `reading_source == "package"` and
   are EXCLUDED from scale-reading accuracy, because the question there is what the
   DISPLAY said and these rows do not answer it. They are not discarded: for Check 1 and
   A23 the package volume may be the better datum, and excluding them from one score is
   not the same as ignoring them.

USAGE
    python parse_validation_workbook.py \
        --xlsx ../outputs/qc/human_validation_v1.xlsx \
        --key  ../outputs/qc/human_validation_v1_key.csv \
        --out  ../outputs/qc/human_verdicts_v1.csv
"""

from __future__ import annotations

import argparse
import os
import re
import sys

import pandas as pd
from openpyxl import load_workbook

# A reading that names a volume unit is a package label, not a scale display. These
# scales print no unit at all on the rows we read, so a unit in the answer can only have
# come from packaging.
UNIT_RE = re.compile(r"\b(m\s?l|ml|millilitre|milliliter|litre|liter|l|g|kg|oz|lb)\b",
                     re.IGNORECASE)
NUM_RE = re.compile(r"[-+]?\d*\.?\d+")
GLARE_RE = re.compile(r"glare|blur|dark|reflect|just about|not sure|unsure|hard to",
                      re.IGNORECASE)


def parse_answer(v):
    """Return (raw_string, numeric_value, unit_named, exact_available)."""
    if v is None:
        return "", None, None, False
    if isinstance(v, (int, float)):
        # Excel coerced it. The number is intact; the literal characters are not.
        return repr(v), float(v), None, False
    s = str(v).strip()
    if not s:
        return "", None, None, False
    unit = None
    m = UNIT_RE.search(s)
    if m:
        unit = m.group(0).replace(" ", "").lower()
    n = NUM_RE.search(s.replace(",", ""))
    val = float(n.group(0)) if n else None
    # A string answer preserves what was typed, so an exact comparison is meaningful.
    return s, val, unit, True


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--key", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--sheet", default="read the display")
    args = ap.parse_args(argv)

    wb = load_workbook(args.xlsx, data_only=True)
    if args.sheet not in wb.sheetnames:
        print(f"sheet '{args.sheet}' not in {wb.sheetnames}", file=sys.stderr)
        return 2
    ws = wb[args.sheet]

    hdr = [c.value for c in ws[1]]
    need = {"image_code", "display_text", "legible", "scale_present", "notes"}
    missing = need - set(hdr)
    if missing:
        print(f"workbook is missing columns: {sorted(missing)}", file=sys.stderr)
        return 2
    col = {name: i + 1 for i, name in enumerate(hdr)}

    rows = []
    for r in range(2, ws.max_row + 1):
        code = ws.cell(row=r, column=col["image_code"]).value
        if code in (None, ""):
            continue
        raw, val, unit, exact = parse_answer(
            ws.cell(row=r, column=col["display_text"]).value)
        leg = str(ws.cell(row=r, column=col["legible"]).value or "").strip()
        scp = str(ws.cell(row=r, column=col["scale_present"]).value or "").strip()
        notes = ws.cell(row=r, column=col["notes"]).value or ""

        blob = f"{leg} {notes}"
        rows.append({
            "image_code": str(code).strip(),
            "h_display_raw": raw,
            "h_display_val": val,
            "h_unit_named": unit,
            "h_exact_available": int(exact),
            # A package volume answers a different question from a scale display.
            "reading_source": "package" if unit in ("ml", "l", "litre", "liter") else "scale",
            "h_legible_raw": leg,
            "h_legible_yes": int(leg.upper().startswith("Y")),
            "h_glare_flag": int(bool(GLARE_RE.search(blob))),
            "h_scale_present": int(scp.lower().startswith("y")),
            "h_notes": str(notes).strip(),
        })

    df = pd.DataFrame(rows)

    key = pd.read_csv(args.key, dtype={"image_code": str})
    df = df.merge(key[["image_code", "id", "filename"]], on="image_code", how="left")

    unmatched = df.id.isna().sum()
    if unmatched:
        print(f"WARNING {unmatched} answered row(s) match no image code in the key",
              file=sys.stderr)
        print(df[df.id.isna()].image_code.tolist(), file=sys.stderr)

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    df.to_csv(args.out, index=False)

    print(f"answered rows              {len(df)}")
    print(f"matched to an id           {int(df.id.notna().sum())}")
    print(f"scale present (human)      {int(df.h_scale_present.sum())}")
    print(f"legible (human)            {int(df.h_legible_yes.sum())}")
    print(f"flagged glare/uncertainty  {int(df.h_glare_flag.sum())}")
    print()
    print("reading_source:")
    print(df.reading_source.value_counts().to_string())
    print()
    print(f"exact-string comparison available on {int(df.h_exact_available.sum())} "
          f"of {len(df)} rows")
    print("  (Excel coerced the rest to numbers, dropping trailing zeros --")
    print("   the numeric comparison is unaffected; see the header.)")
    print(f"\nwrote {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
